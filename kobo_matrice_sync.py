#!/usr/bin/env python3
"""
Synchronisation KoboToolbox -> MATRICE DE SUIVI DES ACTIVITES -> envoi par e-mail.

Étapes :
  1. Récupère les soumissions du formulaire via l'API KoboToolbox (v2).
  2. Garde, pour chaque lot, la visite la plus récente.
  3. Met à jour la dernière matrice trouvée dans MATRIX_DIR et l'enregistre sous
     MATRICE_DE_SUIVI_DES_ACTIVITES_du_JJ-MM-AA.xlsx.
  4. Envoie la matrice par e-mail avec un résumé de la semaine.

Usage :
  python kobo_matrice_sync.py                      # cycle complet (API + mail)
  python kobo_matrice_sync.py --no-mail            # sans envoi d'e-mail
  python kobo_matrice_sync.py --from-json data.json --no-mail   # test hors ligne
"""
import argparse
import datetime as dt
import json
import logging
import os
import re
import smtplib
import ssl
import sys
from copy import copy
from email.message import EmailMessage
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

# --------------------------------------------------------------------------
# Libellés : valeurs Kobo (noms de choix) -> texte écrit dans la matrice
# --------------------------------------------------------------------------
STATUT = {
    "non_demarre": "Non démarré",
    "demarrage_autorise": "Démarrage autorisé",
    "os_a_reprendre": "Ordre de service à reprendre (erreur sur l'OS)",
    "demarrage_effectif": "Démarrage effectif",
    "travaux_en_cours": "Travaux en cours",
    "travaux_acheves": "Travaux achevés",
    "suspendu": "Travaux suspendus",
    "abandonne": "Abandonné",
}
PGES = {
    "valide": "Validé",
    "transmis_validation": "Transmis en cours de validation",
    "non_transmis": "Non transmis",
    "pas_de_pges": "Pas de PGES",
}
ONG = {
    "ong_epc": "ONG EPC",
    "agir_odd": "AGIR pour les ODD",
    "kera_sy_ire": "KERA-SY-IRÊ",
    "etoile_bounkani": "ONG Étoile du Bounkani",
    "cibes": "CIBES",
    "voix_femme": "ONG Voix de Femme",
    "pietro_bonili": "Centre de santé PIETRO BONILI",
    "solidarite": "Solidarité",
    "victoire_plus": "ONG Victoire Plus",
    "esperance_vie": "Espérance Vie",
    "non_mobilisee": "Non mobilisée",
}
STATUTS_A_SURVEILLER = {"suspendu", "abandonne", "os_a_reprendre"}

# Début des intitulés de colonnes de la matrice (comparaison en majuscules sans accents)
HEADERS = {
    "region": "REGION",
    "departement": "DEPARTEMENT",
    "entreprise": "ENTREPRISE",
    "remise": "DATE DE REMISE",
    "taux": "TAUX D",
    "obs_tech": "OBSERVATIONS TECHNIQUES",
    "pges": "OBSERVATIONS ENVIRONNEMENTALES",
    "ong": "ONG MOBILISEES",
    "recep_prov": "DATE DE RECEPTION PROVISOIRE",
    "recep_def": "DATE DE RECEPTION DEFINITIVE",
}
DATE_FMT = "DD/MM/YYYY"
SHEET_HISTORY = "HISTORIQUE"
MATRIX_PATTERN = "MATRICE_DE_SUIVI_DES_ACTIVITES_du_*.xlsx"

log = logging.getLogger("kobo_sync")


# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
def load_env(path: Path) -> None:
    """Charge un fichier .env simple (CLE=valeur) sans écraser l'environnement."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def cfg(name: str, default: str = "", required: bool = False) -> str:
    val = os.environ.get(name, default)
    if required and not val:
        sys.exit(f"Variable de configuration manquante : {name} (voir .env.example)")
    return val


# --------------------------------------------------------------------------
# 1. Kobo
# --------------------------------------------------------------------------
def fetch_submissions() -> list[dict]:
    base = cfg("KOBO_URL", "https://kf.kobotoolbox.org").rstrip("/")
    uid = cfg("KOBO_ASSET_UID", required=True)
    headers = {"Authorization": f"Token {cfg('KOBO_TOKEN', required=True)}"}
    url = f"{base}/api/v2/assets/{uid}/data/?format=json&limit=1000"
    out = []
    while url:
        r = requests.get(url, headers=headers, timeout=60)
        r.raise_for_status()
        payload = r.json()
        out.extend(payload.get("results", []))
        url = payload.get("next")
    log.info("%d soumission(s) récupérée(s) depuis Kobo", len(out))
    return out


def flatten(sub: dict) -> dict:
    """Retire les préfixes de groupes ('avancement/taux_execution' -> 'taux_execution')."""
    flat = {k.split("/")[-1]: v for k, v in sub.items()}
    flat["_id"] = sub.get("_id")
    flat["_submission_time"] = sub.get("_submission_time")
    flat["_validation_status"] = sub.get("_validation_status")
    return flat


def pdate(value) -> dt.date | None:
    if not value:
        return None
    try:
        return dt.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def is_rejected(sub: dict) -> bool:
    vs = sub.get("_validation_status") or {}
    uid = vs.get("uid") if isinstance(vs, dict) else vs
    return uid == "validation_status_not_approved"


def lot_number(sub: dict) -> int | None:
    try:
        return int(str(sub.get("id_projet", "")).strip())
    except ValueError:
        return None


def latest_per_lot(subs: list[dict]) -> dict[int, dict]:
    """Dernière visite (date de visite, puis heure de soumission) par lot."""
    best: dict[int, dict] = {}
    for s in subs:
        if is_rejected(s):
            continue
        n = lot_number(s)
        if n is None:
            log.warning("Soumission %s ignorée : id_projet invalide", s.get("_id"))
            continue
        key = (pdate(s.get("date_mission")) or dt.date.min, str(s.get("_submission_time") or ""))
        if n not in best or key > best[n]["_key"]:
            best[n] = {**s, "_key": key}
    return best


# --------------------------------------------------------------------------
# 2. Matrice
# --------------------------------------------------------------------------
def norm(s) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip().upper()
    return s.translate(str.maketrans("ÉÈÊÀÂÎÔÛÇ", "EEEAAIOUC"))


def find_columns(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, 11):
        found = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(r, c).value)
            for key, prefix in HEADERS.items():
                if key not in found and h.startswith(prefix):
                    found[key] = c
        if "taux" in found and "region" in found:
            return r, found
    sys.exit("En-têtes de la matrice introuvables (ligne 'TAUX D'EXECUTION').")


def lot_rows(ws, header_row: int, cols: dict[str, int]) -> dict[int, int]:
    num_col = cols["region"] - 1  # la colonne des numéros est juste avant 'REGIONS'
    rows = {}
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, num_col).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            rows[int(v)] = r
        elif rows:
            break  # fin du tableau
    return rows


def lot_identity(ws, rows: dict[int, int], cols: dict[str, int]) -> dict[int, dict]:
    ident, region = {}, ""
    for n, r in sorted(rows.items()):
        region = ws.cell(r, cols["region"]).value or region  # cellules fusionnées
        ident[n] = {
            "region": str(region).strip(),
            "departement": str(ws.cell(r, cols["departement"]).value or "").strip(),
            "entreprise": str(ws.cell(r, cols["entreprise"]).value or "").strip(),
        }
    return ident


def set_date(cell, d: dt.date) -> None:
    cell.value = d
    cell.number_format = DATE_FMT


def normalize_taux_scale(ws, rows, col) -> int:
    """Ramène les taux saisis en % (0,05 affiché 5 %) sur l'échelle 0-100 de la matrice."""
    changed = 0
    for r in rows.values():
        c = ws.cell(r, col)
        if "%" in str(c.number_format):
            if isinstance(c.value, (int, float)):
                c.value = round(c.value * 100, 2)
            c.number_format = "General"
            changed += 1
    return changed


def add_trace_columns(ws, header_row, last_col, first_row, last_row) -> tuple[int, int] | None:
    c1, c2 = last_col + 1, last_col + 2
    if any(ws.cell(r, c).value not in (None, "") for r in range(header_row, last_row + 1) for c in (c1, c2)):
        log.warning("Colonnes de traçabilité non ajoutées : cellules à droite déjà utilisées.")
        return None
    model = ws.cell(header_row, last_col)
    for c, title in ((c1, "DATE DERNIERE VISITE"), (c2, "RAPPORTEUR")):
        h = ws.cell(header_row, c, title)
        h.font, h.fill, h.border = copy(model.font), copy(model.fill), copy(model.border)
        h.alignment = copy(model.alignment)
        ws.column_dimensions[h.column_letter].width = 22
    return c1, c2


def update_matrix(wb, latest: dict[int, dict]) -> dict:
    ws = wb["MATRICE DE SUIVI"] if "MATRICE DE SUIVI" in wb.sheetnames else wb.active
    header_row, cols = find_columns(ws)
    rows = lot_rows(ws, header_row, cols)
    ident = lot_identity(ws, rows, cols)
    first_row, last_row = min(rows.values()), max(rows.values())

    n_fixed = normalize_taux_scale(ws, rows, cols["taux"])
    if n_fixed:
        log.info("Échelle des taux harmonisée (0-100) sur %d cellule(s)", n_fixed)

    last_col = max(cols["recep_def"], cols["ong"], cols["pges"])
    trace = add_trace_columns(ws, header_row, last_col, first_row, last_row)

    for n, s in latest.items():
        r = rows.get(n)
        if r is None:
            log.warning("Lot %s absent de la matrice, soumission ignorée", n)
            continue

        statut = s.get("statut")
        obs = (s.get("obs_techniques") or "").strip()
        if statut:
            txt = STATUT.get(statut, statut)
            ws.cell(r, cols["obs_tech"]).value = f"{txt} – {obs}" if obs else txt

        taux = s.get("taux_execution")
        if taux not in (None, ""):
            ws.cell(r, cols["taux"]).value = int(float(taux))
            ws.cell(r, cols["taux"]).number_format = "General"

        d = pdate(s.get("date_remise_site"))
        if d:
            set_date(ws.cell(r, cols["remise"]), d)

        if s.get("pges_statut"):
            ws.cell(r, cols["pges"]).value = PGES.get(s["pges_statut"], s["pges_statut"])

        ong = s.get("ong_mobilisee")
        if ong:
            if ong == "mobilisee_autre":
                nom = (s.get("ong_autre_nom") or "").strip()
                ws.cell(r, cols["ong"]).value = nom or "ONG mobilisée"
            else:
                ws.cell(r, cols["ong"]).value = ONG.get(ong, ong)

        if s.get("recep_prov") == "oui" and pdate(s.get("date_recep_prov")):
            set_date(ws.cell(r, cols["recep_prov"]), pdate(s["date_recep_prov"]))
        if s.get("recep_def") == "oui" and pdate(s.get("date_recep_def")):
            set_date(ws.cell(r, cols["recep_def"]), pdate(s["date_recep_def"]))

        if trace:
            set_date(ws.cell(r, trace[0]), pdate(s.get("date_mission")) or dt.date.today())
            who = s.get("nom_agent") or s.get("agent") or ""
            ws.cell(r, trace[1]).value = str(who).replace("_", " ").title() if who else ""

    # Ligne de total : on conserve la convention du fichier (valeur saisie, pas formule)
    vals = [ws.cell(r, cols["taux"]).value for r in rows.values()]
    vals = [v for v in vals if isinstance(v, (int, float))]
    mean = round(sum(vals) / len(vals), 2) if vals else 0
    for r in range(last_row + 1, last_row + 6):
        label = norm(ws.cell(r, cols["region"]).value)
        if label.startswith("TAUX D"):
            ws.cell(r, cols["taux"]).value = mean
            break

    return {"ws": ws, "rows": rows, "ident": ident, "mean": mean}


def append_history(wb, latest_all: list[dict], ident: dict[int, dict]) -> int:
    """Journal de toutes les visites (une ligne par soumission, sans doublon)."""
    head = ["ID KOBO", "DATE VISITE", "N° LOT", "REGION", "DEPARTEMENT", "ENTREPRISE",
            "STATUT", "TAUX (%)", "PGES", "ONG", "RAPPORTEUR", "SOUMIS LE"]
    if SHEET_HISTORY not in wb.sheetnames:
        hs = wb.create_sheet(SHEET_HISTORY)
        hs.append(head)
        for c in hs[1]:
            c.font = Font(name="Arial", bold=True)
            c.alignment = Alignment(wrap_text=True)
        hs.freeze_panes = "A2"
        for i, w in enumerate([10, 12, 8, 16, 20, 30, 28, 10, 28, 28, 22, 18], 1):
            hs.column_dimensions[hs.cell(1, i).column_letter].width = w
    hs = wb[SHEET_HISTORY]
    known = {str(hs.cell(r, 1).value) for r in range(2, hs.max_row + 1)}
    added = 0
    for s in sorted(latest_all, key=lambda x: (str(x.get("date_mission")), str(x.get("_submission_time")))):
        if is_rejected(s) or str(s.get("_id")) in known:
            continue
        n = lot_number(s)
        i = ident.get(n, {})
        ong = s.get("ong_mobilisee")
        ong_txt = (s.get("ong_autre_nom") or "ONG mobilisée") if ong == "mobilisee_autre" else ONG.get(ong, ong or "")
        hs.append([
            s.get("_id"), pdate(s.get("date_mission")), n, i.get("region", ""), i.get("departement", ""),
            i.get("entreprise", ""), STATUT.get(s.get("statut"), s.get("statut") or ""),
            int(float(s["taux_execution"])) if s.get("taux_execution") not in (None, "") else None,
            PGES.get(s.get("pges_statut"), s.get("pges_statut") or ""), ong_txt,
            s.get("nom_agent") or s.get("agent") or "", str(s.get("_submission_time") or "")[:16].replace("T", " "),
        ])
        hs.cell(hs.max_row, 2).number_format = DATE_FMT
        added += 1
    return added


def find_base_matrix(matrix_dir: Path, explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    files = sorted(matrix_dir.glob(MATRIX_PATTERN), key=matrix_date)
    if not files:
        sys.exit(f"Aucune matrice '{MATRIX_PATTERN}' trouvée dans {matrix_dir}")
    return files[-1]


def matrix_date(p: Path) -> tuple:
    m = re.search(r"du_(\d{2})-(\d{2})-(\d{2})", p.name)
    d = dt.date(2000 + int(m[3]), int(m[2]), int(m[1])) if m else dt.date.min
    return (d, p.stat().st_mtime)


def prune_old(matrix_dir: Path, keep: int, current: Path) -> None:
    """Ne conserve que les `keep` matrices les plus récentes (le dépôt Git garde l'historique)."""
    files = sorted(matrix_dir.glob(MATRIX_PATTERN), key=matrix_date, reverse=True)
    for f in files[keep:]:
        if f != current:
            f.unlink()
            log.info("Ancienne matrice supprimée : %s", f.name)


# --------------------------------------------------------------------------
# 3. Résumé et e-mail
# --------------------------------------------------------------------------
def build_summary(info: dict, latest: dict[int, dict], today: dt.date) -> dict:
    since = today - dt.timedelta(days=7)
    reported, missing, watch = [], [], []
    for n in sorted(info["rows"]):
        s = latest.get(n)
        label = f"Lot {n:02d} – {info['ident'][n]['departement']} ({info['ident'][n]['entreprise']})"
        d = pdate(s.get("date_mission")) if s else None
        if s and d and d >= since:
            reported.append(label)
        else:
            missing.append(label)
        if s and s.get("statut") in STATUTS_A_SURVEILLER:
            watch.append(f"{label} : {STATUT[s['statut']]}")
    return {"reported": reported, "missing": missing, "watch": watch, "mean": info["mean"], "total": len(info["rows"])}


def mail_body(summ: dict, today: dt.date) -> str:
    lines = [
        "Bonjour,",
        "",
        f"Veuillez trouver ci-joint la matrice de suivi des activités mise à jour au {today:%d/%m/%Y}.",
        "",
        f"Taux d'exécution moyen : {summ['mean']} %",
        f"Lots renseignés ces 7 derniers jours : {len(summ['reported'])} sur {summ['total']}",
    ]
    if summ["missing"]:
        lines += ["", "Lots sans visite cette semaine :"] + [f"  - {x}" for x in summ["missing"]]
    if summ["watch"]:
        lines += ["", "Points d'attention :"] + [f"  - {x}" for x in summ["watch"]]
    lines += ["", "Message généré automatiquement à partir des données KoboToolbox.", "", "Cordialement"]
    return "\n".join(lines)


def send_mail(attachment: Path, summ: dict, today: dt.date) -> None:
    host = cfg("SMTP_HOST", required=True)
    port = int(cfg("SMTP_PORT", "587"))
    user = cfg("SMTP_USER", required=True)
    password = cfg("SMTP_PASSWORD", required=True)
    sender = cfg("MAIL_FROM", user)
    to = [a.strip() for a in cfg("MAIL_TO", required=True).split(",") if a.strip()]

    msg = EmailMessage()
    msg["Subject"] = f"Matrice de suivi des activités – {today:%d/%m/%Y}"
    msg["From"], msg["To"] = sender, ", ".join(to)
    msg.set_content(mail_body(summ, today))
    msg.add_attachment(
        attachment.read_bytes(), maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=attachment.name,
    )
    ctx = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctx, timeout=60) as smtp:
            smtp.login(user, password)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=60) as smtp:
            smtp.starttls(context=ctx)
            smtp.login(user, password)
            smtp.send_message(msg)
    log.info("E-mail envoyé à %s", ", ".join(to))


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--from-json", help="Fichier JSON de soumissions (test hors ligne)")
    ap.add_argument("--base", help="Matrice de départ (par défaut : la plus récente de MATRIX_DIR)")
    ap.add_argument("--no-mail", action="store_true", help="Ne pas envoyer d'e-mail")
    ap.add_argument("--skip-if-unchanged", action="store_true",
                    help="Ne rien enregistrer s'il n'y a aucune nouvelle soumission")
    ap.add_argument("--env", default=".env", help="Fichier de configuration (défaut : .env)")
    args = ap.parse_args()

    load_env(Path(args.env))
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(), logging.FileHandler("sync.log", encoding="utf-8")])

    matrix_dir = Path(cfg("MATRIX_DIR", "./matrices"))
    matrix_dir.mkdir(parents=True, exist_ok=True)
    today = dt.date.today()

    if args.from_json:
        raw = json.loads(Path(args.from_json).read_text(encoding="utf-8"))
        raw = raw.get("results", raw) if isinstance(raw, dict) else raw
    else:
        raw = fetch_submissions()
    subs = [flatten(s) for s in raw]
    latest = latest_per_lot(subs)
    log.info("%d lot(s) avec au moins une visite valide", len(latest))

    base = find_base_matrix(matrix_dir, args.base)
    log.info("Matrice de départ : %s", base.name)
    wb = load_workbook(base)
    info = update_matrix(wb, latest)
    added = append_history(wb, subs, info["ident"])
    log.info("%d ligne(s) ajoutée(s) à l'onglet %s", added, SHEET_HISTORY)
    if args.skip_if_unchanged and added == 0:
        log.info("Aucune nouvelle soumission : matrice inchangée, rien n'est enregistré.")
        return

    out = matrix_dir / f"MATRICE_DE_SUIVI_DES_ACTIVITES_du_{today:%d-%m-%y}.xlsx"
    wb.save(out)
    log.info("Matrice enregistrée : %s (taux moyen %s %%)", out, info["mean"])
    prune_old(matrix_dir, keep=int(cfg("MATRIX_KEEP", "8")), current=out)
    if os.environ.get("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a", encoding="utf-8") as fh:
            fh.write(f"matrix_path={out}\n")

    if args.no_mail:
        log.info("Envoi d'e-mail désactivé (--no-mail)")
        return
    send_mail(out, build_summary(info, latest, today), today)


if __name__ == "__main__":
    main()
